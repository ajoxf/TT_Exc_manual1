"""Shared styles and helpers for the TT z-score monitor workbook builder."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

# ---- palette -------------------------------------------------------------
C_TITLE      = "1F3864"
C_HDR_FILL   = "1F3864"
C_HDR_TEXT   = "FFFFFF"
C_SECT_FILL  = "D9E1F2"
C_BLOCK_FILL = "EDF2FA"
C_INPUT_FILL = "FFFF00"   # cells the user may edit
C_INPUT_TEXT = "0000FF"   # hardcoded input
C_FORMULA    = "000000"
C_LINK       = "008000"   # link to another sheet
C_NOTE       = "7F7F7F"
C_WARN       = "C00000"
C_GREEN      = "C6EFCE"; C_GREEN_T = "006100"
C_AMBER      = "FFEB9C"; C_AMBER_T = "9C5700"
C_RED        = "FFC7CE"; C_RED_T   = "9C0006"
C_GREY       = "E7E6E6"; C_GREY_T  = "595959"

THIN = Side(style="thin", color="BFBFBF")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def f(size=10, bold=False, italic=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)

def title(ws, cell, text, size=16):
    ws[cell] = text
    ws[cell].font = f(size, bold=True, color=C_TITLE)

def block_title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = f(11, bold=True, color=C_TITLE)

def note(ws, cell, text, color=C_NOTE, italic=True, bold=False):
    ws[cell] = text
    ws[cell].font = f(9, bold=bold, italic=italic, color=color)

def header_row(ws, row, cols, fill=C_HDR_FILL, text=C_HDR_TEXT):
    """cols: dict of column-letter -> header text"""
    for col, txt in cols.items():
        c = ws[f"{col}{row}"]
        c.value = txt
        c.font = Font(name=FONT, size=9, bold=True, color=text)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX

def label(ws, cell, text, bold=False, size=10, color="000000", indent=0):
    ws[cell] = text
    ws[cell].font = f(size, bold=bold, color=color)
    if indent:
        ws[cell].alignment = Alignment(indent=indent)

def input_cell(ws, cell, value, numfmt=None):
    c = ws[cell]
    c.value = value
    c.font = f(10, bold=True, color=C_INPUT_TEXT)
    c.fill = PatternFill("solid", fgColor=C_INPUT_FILL)
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if numfmt:
        c.number_format = numfmt
    return c

def value_cell(ws, cell, value=None, numfmt=None, bold=False, align="center", fill=None):
    c = ws[cell]
    c.value = value
    c.font = f(10, bold=bold)
    c.border = BOX
    c.alignment = Alignment(horizontal=align)
    if numfmt:
        c.number_format = numfmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c

def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def sect(ws, row, text, first="B", last="G"):
    """Shaded section divider across a row."""
    ws[f"{first}{row}"] = text
    ws[f"{first}{row}"].font = f(9, bold=True, color=C_TITLE)
    for col in range(ord(first), ord(last) + 1):
        ws[f"{chr(col)}{row}"].fill = PatternFill("solid", fgColor=C_SECT_FILL)
