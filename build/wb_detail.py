"""Detail: everything the main Dashboard no longer shows.

Full per-spread breakdown for all four slots - touch, window, warm-up gates,
feed rate, frozen statistics, costs, break-even, take-profit and the edge
filter. Also carries the outright legs and the derived/listed spread table.

No formulas here either; VBA writes every value, compare-first.
"""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from wb_common import *

PX = "#,##0.0000"
USD = '$#,##0.00;($#,##0.00);-'
Z2 = "0.00"

LEGS = [
    (6, "RBV6", "RBOB Gasoline Oct26", "42,000 gal"),
    (7, "HOV6", "NY Harbor ULSD Oct26", "42,000 gal"),
    (8, "CLV6", "WTI Crude Oct26", "1,000 bbl"),
    (9, "BZV6", "Brent Last Day Fin Oct26", "1,000 bbl"),
]

DERIVED = [
    (13, "BZ - CL", "legged", "bid = BZ bid - CL ask ; ask = BZ ask - CL bid"),
    (14, "HO/CL crack", "legged", "bid = HO bid x42 - CL ask ; ask = HO ask x42 - CL bid"),
    (15, "3:2:1 crack", "legged", "bid sells products at the bid and buys crude at the ask; ask mirrors it"),
    (16, "CL Oct26 - BZ Oct26 Inter-Product", "listed", "exchange-listed - one quote, no leg risk"),
    (17, "Oct26 HO-CL Crack", "listed", "exchange-listed - one quote, no leg risk"),
]

# (row, label, unit/note, kind)
METRICS = [
    (21, "Spread", "", "text"),
    (22, "Definition", "LegB - beta x LegA", "text"),
    (23, "Mode", "", "text"),
    (24, "HEDGE_RATIO (beta)", "stamped pair ->", "text"),
    (25, "USD per 1.00 of spread", "per lot", "usd0"),
    (26, "SECT:MARKET  (touch, shown separately from the statistic)", "", "sect"),
    (27, "Spread bid", "", "px"),
    (28, "Spread ask", "", "px"),
    (29, "Spread mid   <- feeds the statistic", "", "px"),
    (30, "Bid-Ask Gap", "spread units", "px"),
    (31, "SECT:WINDOW", "", "sect"),
    (32, "Stored samples in window", "after dedupe + store gate", "int"),
    (33, "Window elapsed", "min", "num2"),
    (34, "Warm-up gate", "binding gate", "text"),
    (35, "Feed rate", "quote CHANGES / min", "num1"),
    (36, "Feed status", "", "text"),
    (37, "SECT:STATISTICS  (FROZEN - republished every STATS_REFRESH_MIN)", "", "sect"),
    (38, "Mean (frozen)", "", "px"),
    (39, "Std Dev (frozen)", "", "px"),
    (40, "Std Dev in dollars", "sigma x USD/pt x LOTS", "usd"),
    (41, "Stats age / next refresh", "min", "text"),
    (42, "Z  (live, vs frozen mean & sigma)", "", "z"),
    (43, "SIGNAL", "", "text"),
    (44, "SECT:COST OF THE ROUND TRIP", "", "sect"),
    (45, "Direction", "", "text"),
    (46, "Entry reference (touch)", "", "px"),
    (47, "Commission", "contracts x rate x LOTS", "usd"),
    (48, "Crossing cost - LEGGING", "", "usd"),
    (49, "Crossing cost - LISTED", "", "usd"),
    (50, "TOTAL COST (this mode)", "", "usd"),
    (51, "SECT:TAKE PROFIT  (exits act on MONEY, never on z)", "", "sect"),
    (52, "TP rule in force", "", "text"),
    (53, "Notional (per NOTIONAL_BASIS)", "", "usd"),
    (54, "Win target (WIN_PCT x notional)", "", "usd"),
    (55, "Costs in spread units", "", "px"),
    (56, "Win in spread units", "", "px"),
    (57, "Break-even spread (mid)", "absolute level, mid scale", "px"),
    (58, "TAKE-PROFIT spread (mid)", "absolute level, mid scale", "px"),
    (59, "TP distance", "required MID move", "px"),
    (60, "TP in sigma", "x sigma", "num2"),
    (61, "Target z", "z at the TP level", "z"),
    (62, "TP beyond the mean?", "", "text"),
    (63, "SECT:EDGE FILTER  (a separate gate from ENTRY_Z)", "", "sect"),
    (64, "Expected capture", "0.5 x |z| x sigma x USD/pt x LOTS", "usd"),
    (65, "Required", "EDGE_MULTIPLE x total cost", "usd"),
    (66, "EDGE VERDICT", "", "text"),
    (67, "Min sigma to pass at live z", "spread units", "px"),
    (69, "SECT:THE MEAN AT THE TOUCH  (live Gap - what you would actually deal at)", "", "sect"),
    (70, "Mean for SHORTING the spread", "mean - Gap/2   (you sell the bid)", "px"),
    (71, "Mean for LONGING the spread", "mean + Gap/2   (you buy the ask)", "px"),
]

SPREAD_COLS = ["D", "E", "F", "G"]


def build(wb):
    ws = wb.create_sheet("Detail")
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 3, "B": 32, "C": 20, "D": 18, "E": 18, "F": 18, "G": 18,
                    "H": 12, "I": 12, "J": 12, "K": 12, "L": 14, "M": 46})
    ws.freeze_panes = "B5"

    title(ws, "B1", "DETAIL   -   full breakdown behind the Dashboard", 14)
    note(ws, "B2", "Same engine, same numbers. The Dashboard shows the three spreads you trade from; "
                   "this sheet shows everything that stands behind them, for all four slots.")

    block_title(ws, "B4", "OUTRIGHT LEGS")
    header_row(ws, 5, {"B": "Symbol", "C": "Instrument ID", "D": "Description", "E": "Contract",
                       "H": "Bid", "I": "Ask", "J": "High", "K": "Low", "L": "Mid", "M": "Gap"})
    for row, sym, desc, contract in LEGS:
        value_cell(ws, f"B{row}", sym, bold=True, align="left")
        value_cell(ws, f"C{row}", "", align="left"); ws[f"C{row}"].font = f(9, color=C_LINK)
        value_cell(ws, f"D{row}", desc, align="left"); ws[f"D{row}"].font = f(9)
        value_cell(ws, f"E{row}", contract); ws[f"E{row}"].font = f(9, color=C_NOTE)
        for col in "HIJKLM":
            value_cell(ws, f"{col}{row}", None, PX)

    block_title(ws, "B11", "DERIVED & LISTED SPREADS")
    header_row(ws, 12, {"B": "Spread", "C": "Basis", "D": "Crossing convention",
                        "H": "Bid", "I": "Ask", "L": "Mid", "M": "Gap"})
    for row, name, basis, conv in DERIVED:
        value_cell(ws, f"B{row}", name, bold=True, align="left")
        value_cell(ws, f"C{row}", basis); ws[f"C{row}"].font = f(9, color=C_NOTE)
        value_cell(ws, f"D{row}", conv, align="left"); ws[f"D{row}"].font = f(9, color=C_NOTE)
        for col in ("H", "I", "L", "M"):
            value_cell(ws, f"{col}{row}", None, PX)

    block_title(ws, "B19", "PER-SPREAD BREAKDOWN")
    hdr = {"B": "Metric", "C": "Unit / note"}
    for i, col in enumerate(SPREAD_COLS, start=1):
        hdr[col] = f"SPREAD {i}"
    header_row(ws, 20, hdr)

    fmt_map = {"px": PX, "usd": USD, "usd0": "#,##0", "num1": "0.0", "num2": "0.00",
               "int": "#,##0", "z": Z2, "text": None}
    bold_rows = {"SIGNAL", "EDGE VERDICT", "TOTAL COST (this mode)", "TAKE-PROFIT spread (mid)",
                 "Mean for SHORTING the spread", "Mean for LONGING the spread",
                 "Z  (live, vs frozen mean & sigma)"}

    for row, lab, unit, kind in METRICS:
        if kind == "sect":
            sect(ws, row, lab.replace("SECT:", ""), first="B", last="G")
            continue
        label(ws, f"B{row}", lab, bold=(lab in bold_rows))
        ws[f"B{row}"].border = BOX
        value_cell(ws, f"C{row}", unit, align="left"); ws[f"C{row}"].font = f(9, color=C_NOTE)
        for col in SPREAD_COLS:
            c = value_cell(ws, f"{col}{row}", None, fmt_map[kind], bold=(lab in bold_rows))
            if kind == "text":
                c.alignment = Alignment(horizontal="center")

    def rule(fill, text_color, formula):
        return FormulaRule(formula=[formula], stopIfTrue=False,
                           fill=PatternFill("solid", bgColor=fill),
                           font=Font(name=FONT, size=10, bold=True, color=text_color))

    for r_ in (
        rule(C_RED,   C_RED_T,   '=OR(D$43="NO USABLE Z",D$43="STALE FEED")'),
        rule(C_GREY,  C_GREY_T,  '=D$43="WARMING UP"'),
        rule(C_AMBER, C_AMBER_T, '=D$43="ABOVE CEILING"'),
        rule(C_GREEN, C_GREEN_T, '=D$43="ENTRY BAND"'),
    ):
        ws.conditional_formatting.add("D42:G43", r_)
    ws.conditional_formatting.add("D66:G66", rule(C_GREEN, C_GREEN_T, '=D66="PASS"'))
    ws.conditional_formatting.add("D66:G66", rule(C_AMBER, C_AMBER_T, '=D66="FAIL"'))
    ws.conditional_formatting.add("D36:G36", rule(C_RED,   C_RED_T,   '=D36="STALE"'))
    ws.conditional_formatting.add("D36:G36", rule(C_AMBER, C_AMBER_T, '=D36="THIN"'))
    ws.conditional_formatting.add("D36:G36", rule(C_GREEN, C_GREEN_T, '=D36="OK"'))
    ws.conditional_formatting.add("D34:G34", rule(C_GREY,  C_GREY_T,  '=D34<>"ready"'))
    ws.conditional_formatting.add("D62:G62", rule(C_AMBER, C_AMBER_T, '=LEFT(D62,3)="YES"'))

    note(ws, "B73", "The EDGE FILTER is a different gate from ENTRY_Z. ENTRY_Z decides when the chime fires; "
                    "the edge filter asks whether the expected capture clears the round trip by EDGE_MULTIPLE. "
                    "A highlighted z that fails the edge filter is still not a trade worth taking.",
         color=C_WARN, bold=True)
    note(ws, "B75", "Break-even and take-profit are MID levels, on the same scale as the mean, sigma and z - "
                    "a mid move of the cost is exactly what pays the round trip. The entry reference above "
                    "them is the TOUCH you actually get filled at. To work the exit order, cross back: add "
                    "half the Bid-Ask Gap to close a short, subtract half to close a long.")
    note(ws, "B74", "ALERT_ONLY_IF_EDGE_PASSES on Config decides whether a failing edge filter also silences the "
                    "chime. It ships FALSE, so the chime follows ENTRY_Z alone.")
    return ws
