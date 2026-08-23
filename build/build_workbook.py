"""Assemble TT_ZScore_Monitor.xlsx.

Two phases, because LibreOffice cannot evaluate RTD():
  phase 1  everything except the RTD formulas  -> recalc.py validates the real formulas
  phase 2  inject the Feed RTD formulas        -> saved without recalculation
Run:  python3 build_workbook.py phase1 <out>   /   python3 build_workbook.py phase2 <out>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook, load_workbook
import wb_common, wb_config, wb_dashboard, wb_detail, wb_feed, wb_other


PARAM_NAMES = {n for grp in (wb_config.ENGINE, wb_config.WINDOW, wb_config.SIGNAL,
                             wb_config.COSTS, wb_config.ALERTS, wb_config.PERSIST)
               for n, *_ in grp} | {"COMMISSION_CONSERVATIVE", "COMMISSION_IF_CME_PER_RT",
                                    "COMMISSION_BASIS", "COMMISSION_PER_LOT_ROUND_TURN"}


def cfg_row_map(ws):
    m = {}
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if isinstance(v, str) and v in PARAM_NAMES:
            m[v] = row
    missing = PARAM_NAMES - set(m)
    if missing:
        raise SystemExit(f"Config parameters not found on the sheet: {sorted(missing)}")
    return m


def phase1(out):
    wb = Workbook()
    wb.remove(wb.active)

    wb_dashboard.build(wb)
    wb_detail.build(wb)
    cfg, tbl_marker, first_spread_row = wb_config.build(wb)
    wb_feed.build(wb)
    wb_other.build_buffer(wb)
    wb_other.build_log(wb)
    wb_other.build_setup(wb, cfg_row_map(cfg))

    for ws in wb:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = False

    wb["Feed"].sheet_state = "hidden"
    wb["Buffer"].sheet_state = "hidden"
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)

    print(f"phase1 ok  sheets={wb.sheetnames}")
    print(f"config spread table marker row={tbl_marker} first spread row={first_spread_row}")


def normalise_booleans(wb):
    """LibreOffice's round-trip rewrites boolean cells as =TRUE()/=FALSE()
    formulas. They evaluate correctly, but these are cells the user edits, and a
    formula in a yellow input cell is a trap. Put the literals back."""
    n = 0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    v = c.value.strip().upper()
                    if v in ("=TRUE()", "=FALSE()"):
                        c.value = (v == "=TRUE()")
                        n += 1
    return n


def phase2(out):
    wb = load_workbook(out)
    n = normalise_booleans(wb)
    print(f"phase2: {n} boolean cells restored to literals")
    wb_feed.inject_formulas(wb["Feed"])
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("phase2 ok  - RTD formulas injected into Feed")


if __name__ == "__main__":
    mode, out = sys.argv[1], sys.argv[2]
    {"phase1": phase1, "phase2": phase2}[mode](out)
