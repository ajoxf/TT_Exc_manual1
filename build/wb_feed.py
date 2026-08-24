"""Feed (hidden): every RTD() formula in the workbook lives here, and nowhere else.

Layout is fixed - the VBA reads these exact ranges.
  B2                market  (-> Dashboard!C4)
  rows  6..11       instruments: A sym | B instId | C bid | D ask | E high | F low | G mid | H width
  rows 16..20       derived + listed spreads (display): A name | B basis | C bid | D ask | E mid | F width
  rows 25..32       MONITOR LEGS, two rows per slot: A slot | B side | C label | D bid | E ask | F mid
"""
from openpyxl.styles import Alignment, PatternFill
from wb_common import *

INSTRUMENTS = [
    (6,  "RBV6", "RBOB Gasoline Oct26"),
    (7,  "HOV6", "NY Harbor ULSD Oct26"),
    (8,  "CLV6", "WTI Crude Oct26"),
    (9,  "BZV6", "Brent Last Day Financial Oct26"),
    (10, "CL Oct26 - BZ Oct26 Inter-Product", "Exchange-listed inter-product spread"),
    (11, "Oct26 HO-CL Crack", "Exchange-listed crack spread"),
]
R_RB, R_HO, R_CL, R_BZ, R_CLBZ, R_HOCL = 6, 7, 8, 9, 10, 11

DERIVED = [
    (16, "BZ - CL",                            "legged"),
    (17, "HO/CL crack",                        "legged"),
    (18, "3:2:1 crack",                        "legged"),
    (19, "CL Oct26 - BZ Oct26 Inter-Product",  "listed"),
    (20, "Oct26 HO-CL Crack",                  "listed"),
]

# slot -> (legA row, legB row); rows 25..32
LEG_ROWS = {1: (25, 26), 2: (27, 28), 3: (29, 30), 4: (31, 32)}
LEG_LABELS = {
    1: ("CLV6",                                   "BZV6"),
    2: ("CLV6",                                   "HOV6 x 42"),
    3: ("CLV6",                                   "(2 x RBV6 + 1 x HOV6) x 42 / 3"),
    4: ("(none - LISTED spread)",                 "CL Oct26 - BZ Oct26 Inter-Product"),
}
# slot -> Feed derived row used for the LISTED crossing-cost comparison ("" = none)
LISTED_REF = {1: 19, 2: 20, 3: 0, 4: 19}


def build(wb):
    ws = wb.create_sheet("Feed")
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 40, "B": 34, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 60})

    title(ws, "A1", "FEED  -  hidden.  Every RTD() formula lives here.", 13)
    note(ws, "A3", "Never put an RTD() formula on Dashboard: a cell holding one repaints whenever the feed moves, "
                   "and that cannot be prevented while the formula is in the cell.", color=C_WARN, bold=True)
    label(ws, "A2", "Market", bold=True)

    label(ws, "A5", "INSTRUMENTS", bold=True, color=C_TITLE)
    header_row(ws, 5, {"B": "Instrument ID", "C": "Bid", "D": "Ask", "E": "High", "F": "Low",
                       "G": "Mid", "H": "Width"})
    ws["A5"].fill = PatternFill("solid", fgColor=C_HDR_FILL)
    ws["A5"].font = f(9, bold=True, color=C_HDR_TEXT)
    for row, sym, desc in INSTRUMENTS:
        value_cell(ws, f"A{row}", sym, align="left", bold=True)
        label(ws, f"I{row}", desc, size=9, color=C_NOTE)
        for col in "BCDEFGH":
            ws[f"{col}{row}"].border = BOX

    label(ws, "A15", "DERIVED & LISTED SPREADS (display)", bold=True, color=C_TITLE)
    header_row(ws, 15, {"B": "Basis", "C": "Bid", "D": "Ask", "E": "Mid", "F": "Width"})
    ws["A15"].fill = PatternFill("solid", fgColor=C_HDR_FILL)
    ws["A15"].font = f(9, bold=True, color=C_HDR_TEXT)
    for row, name, basis in DERIVED:
        value_cell(ws, f"A{row}", name, align="left", bold=True)
        value_cell(ws, f"B{row}", basis)
        for col in "CDEF":
            ws[f"{col}{row}"].border = BOX

    label(ws, "A23", "MONITOR LEGS", bold=True, color=C_TITLE)
    note(ws, "B23", "Edit these formulas to redefine what a monitored spread is made of. "
                    "spread = LegB - HEDGE_RATIO x LegA, and nothing else - no carry, no fair-value term.")
    header_row(ws, 24, {"A": "Slot", "B": "Side", "C": "Composition", "D": "Bid", "E": "Ask", "F": "Mid"})
    for slot, (ra, rb) in LEG_ROWS.items():
        la, lb = LEG_LABELS[slot]
        for row, side, lab in ((ra, "LegA", la), (rb, "LegB", lb)):
            value_cell(ws, f"A{row}", slot, "0")
            value_cell(ws, f"B{row}", side, bold=True)
            value_cell(ws, f"C{row}", lab, align="left")
            for col in "DEF":
                ws[f"{col}{row}"].border = BOX

    label(ws, "A34", "LISTED-EQUIVALENT ROW PER SLOT  (for the side-by-side crossing-cost comparison)",
          bold=True, color=C_TITLE)
    header_row(ws, 35, {"A": "Slot", "B": "Feed derived row", "C": "Name"})
    for i, slot in enumerate(sorted(LISTED_REF), start=36):
        value_cell(ws, f"A{i}", slot, "0")
        value_cell(ws, f"B{i}", LISTED_REF[slot] or "", "0")
        nm = next((n for r, n, _ in DERIVED if r == LISTED_REF[slot]), "(none)")
        value_cell(ws, f"C{i}", nm, align="left")

    note(ws, "A41", "Crossing convention used throughout:  spread bid = LegB bid - beta x LegA ask ; "
                    "spread ask = LegB ask - beta x LegA bid ; spread mid = LegB mid - beta x LegA mid.")
    note(ws, "A42", "For a LISTED slot LegA is empty and is treated as zero, so the listed quote passes straight "
                    "through untouched.")
    return ws


def inject_formulas(ws, market_ref="=Dashboard!$C$4"):
    """Phase 2 - written AFTER LibreOffice recalculation.

    LibreOffice cannot evaluate RTD(); recalculating with these in place would bake #NAME?
    into the delivered file and lower-case the formulas it failed to parse.
    """
    ws["B2"] = market_ref
    ws["B2"].font = f(10, bold=True, color=C_LINK)

    for row, sym, _ in INSTRUMENTS:
        ws[f"B{row}"] = f'=RTD("tt.rtd",,"Inst",$B$2,$A{row})'
        # Guard on TT's status strings too. When the platform is not logged in
        # the id cell holds "Not_Connected", and asking the RTD server for a
        # Bid on that registers a meaningless topic.
        for col, fld in (("C", "Bid"), ("D", "Ask"), ("E", "High"), ("F", "Low")):
            ws[f"{col}{row}"] = (f'=IFERROR(IF(OR($B{row}="",LEFT($B{row},4)="Not_"),"",'
                                 f'VALUE(RTD("tt.rtd",,$B{row},"{fld}"))),"")')
        ws[f"G{row}"] = f'=IF(OR($C{row}="",$D{row}=""),"",($C{row}+$D{row})/2)'
        ws[f"H{row}"] = f'=IF(OR($C{row}="",$D{row}=""),"",$D{row}-$C{row})'

    def two(a, b):   # "both present" guard
        return f'OR({a}="",{b}="")'

    # BZ - CL   (LegB = BZ, LegA = CL)
    ws[f"C16"] = f'=IF({two(f"$C${R_BZ}", f"$D${R_CL}")},"",$C${R_BZ}-$D${R_CL})'
    ws[f"D16"] = f'=IF({two(f"$D${R_BZ}", f"$C${R_CL}")},"",$D${R_BZ}-$C${R_CL})'
    ws[f"E16"] = f'=IF({two(f"$G${R_BZ}", f"$G${R_CL}")},"",$G${R_BZ}-$G${R_CL})'
    # HO/CL crack
    ws[f"C17"] = f'=IF({two(f"$C${R_HO}", f"$D${R_CL}")},"",$C${R_HO}*42-$D${R_CL})'
    ws[f"D17"] = f'=IF({two(f"$D${R_HO}", f"$C${R_CL}")},"",$D${R_HO}*42-$C${R_CL})'
    ws[f"E17"] = f'=IF({two(f"$G${R_HO}", f"$G${R_CL}")},"",$G${R_HO}*42-$G${R_CL})'
    # 3:2:1 crack - bid sells products at the bid and buys crude at the ASK; ask mirrors it
    ws[f"C18"] = (f'=IF(OR($C${R_RB}="",$C${R_HO}="",$D${R_CL}=""),"",'
                  f'(($C${R_RB}*2*42)+($C${R_HO}*1*42)-($D${R_CL}*3))/3)')
    ws[f"D18"] = (f'=IF(OR($D${R_RB}="",$D${R_HO}="",$C${R_CL}=""),"",'
                  f'(($D${R_RB}*2*42)+($D${R_HO}*1*42)-($C${R_CL}*3))/3)')
    ws[f"E18"] = (f'=IF(OR($G${R_RB}="",$G${R_HO}="",$G${R_CL}=""),"",'
                  f'(($G${R_RB}*2*42)+($G${R_HO}*1*42)-($G${R_CL}*3))/3)')
    # listed spreads pass straight through
    for row, src in ((19, R_CLBZ), (20, R_HOCL)):
        ws[f"C{row}"] = f"=$C${src}"
        ws[f"D{row}"] = f"=$D${src}"
        ws[f"E{row}"] = f"=$G${src}"
    for row, _, _ in DERIVED:
        ws[f"F{row}"] = f'=IF(OR($C{row}="",$D{row}=""),"",$D{row}-$C{row})'

    # ---- monitor legs -----------------------------------------------------
    def leg(row, bid, ask, mid):
        ws[f"D{row}"], ws[f"E{row}"], ws[f"F{row}"] = bid, ask, mid

    leg(25, f"=$C${R_CL}", f"=$D${R_CL}", f"=$G${R_CL}")                      # 1 LegA  CL
    leg(26, f"=$C${R_BZ}", f"=$D${R_BZ}", f"=$G${R_BZ}")                      # 1 LegB  BZ
    leg(27, f"=$C${R_CL}", f"=$D${R_CL}", f"=$G${R_CL}")                      # 2 LegA  CL
    leg(28, f'=IF($C${R_HO}="","",$C${R_HO}*42)',
            f'=IF($D${R_HO}="","",$D${R_HO}*42)',
            f'=IF($G${R_HO}="","",$G${R_HO}*42)')                             # 2 LegB  HO x42
    leg(29, f"=$C${R_CL}", f"=$D${R_CL}", f"=$G${R_CL}")                      # 3 LegA  CL
    leg(30, f'=IF(OR($C${R_RB}="",$C${R_HO}=""),"",($C${R_RB}*2*42+$C${R_HO}*1*42)/3)',
            f'=IF(OR($D${R_RB}="",$D${R_HO}=""),"",($D${R_RB}*2*42+$D${R_HO}*1*42)/3)',
            f'=IF(OR($G${R_RB}="",$G${R_HO}=""),"",($G${R_RB}*2*42+$G${R_HO}*1*42)/3)')  # 3 LegB
    leg(31, "", "", "")                                                       # 4 LegA  (listed)
    leg(32, f"=$C${R_CLBZ}", f"=$D${R_CLBZ}", f"=$G${R_CLBZ}")                # 4 LegB  listed
